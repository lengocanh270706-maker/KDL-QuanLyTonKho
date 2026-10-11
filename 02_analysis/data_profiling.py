from pathlib import Path
import pandas as pd
import numpy as np

BASE_DIR = Path(__file__).resolve().parent
INPUT_FILE = BASE_DIR.parent / "data1.xlsx"
OUTPUT_DIR = BASE_DIR / "output"
OUTPUT_FILE = OUTPUT_DIR / "data_quality_report.xlsx"

def find_outliers(df, sheet_name):
    """Phát hiện outlier số bằng quy tắc IQR; bỏ qua cột số có ít hơn 4 giá trị hợp lệ."""
    rows = []
    numeric_cols = df.select_dtypes(include=[np.number]).columns
    for col in numeric_cols:
        series = pd.to_numeric(df[col], errors="coerce").dropna()
        if len(series) < 4 or series.nunique() < 2:
            continue
        q1 = series.quantile(0.25)
        q3 = series.quantile(0.75)
        iqr = q3 - q1
        if iqr == 0:
            lower, upper = q1, q3
        else:
            lower = q1 - 1.5 * iqr
            upper = q3 + 1.5 * iqr
        mask = pd.to_numeric(df[col], errors="coerce").notna() & (
            (pd.to_numeric(df[col], errors="coerce") < lower) |
            (pd.to_numeric(df[col], errors="coerce") > upper)
        )
        for idx in df.index[mask]:
            rows.append({
                "Sheet": sheet_name,
                "Excel_Row": int(idx) + 2,
                "Column": str(col),
                "Value": df.loc[idx, col],
                "IQR_Lower_Bound": lower,
                "IQR_Upper_Bound": upper,
                "Note": "Giá trị ngoài ngưỡng IQR; cần kiểm tra nghiệp vụ, không tự động kết luận là sai."
            })
    return rows

def main():
    if not INPUT_FILE.exists():
        raise FileNotFoundError(
            f"Không tìm thấy {INPUT_FILE}. Hãy đặt data1.xlsx ở thư mục gốc dự án, "
            "cùng cấp với thư mục 02_analysis."
        )

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    excel = pd.ExcelFile(INPUT_FILE)

    overview_rows, profile_rows, duplicate_rows, outlier_rows, invalid_rows = [], [], [], [], []

    for sheet in excel.sheet_names:
        df = pd.read_excel(INPUT_FILE, sheet_name=sheet)
        row_count, col_count = df.shape
        exact_duplicate_count = int(df.duplicated(keep="first").sum())

        overview_rows.append({
            "Sheet": sheet,
            "Rows": row_count,
            "Columns": col_count,
            "Exact_Duplicate_Rows": exact_duplicate_count,
            "Total_Null_Cells": int(df.isna().sum().sum()),
            "Null_Cell_Rate_%": round(df.isna().sum().sum() / (row_count * col_count) * 100, 2)
                if row_count * col_count else 0
        })

        for col in df.columns:
            s = df[col]
            null_count = int(s.isna().sum())
            non_null = s.dropna()
            profile_rows.append({
                "Sheet": sheet,
                "Column": str(col),
                "Pandas_Dtype": str(s.dtype),
                "Rows": row_count,
                "Non_Null_Count": int(s.notna().sum()),
                "Null_Count": null_count,
                "Null_Rate_%": round(null_count / row_count * 100, 2) if row_count else 0,
                "Distinct_Non_Null": int(s.nunique(dropna=True)),
                "Sample_Values": " | ".join(map(str, non_null.head(3).tolist()))
            })

        # Liệt kê các dòng trùng hoàn toàn (bỏ qua bản ghi xuất hiện đầu tiên)
        dup_mask = df.duplicated(keep="first")
        for idx in df.index[dup_mask]:
            row = {"Sheet": sheet, "Excel_Row": int(idx) + 2}
            row.update({str(c): df.loc[idx, c] for c in df.columns})
            duplicate_rows.append(row)

        outlier_rows.extend(find_outliers(df, sheet))

        # Kiểm tra giá trị âm ở các cột số; đây là cờ cảnh báo chung, cần xác minh nghiệp vụ.
        for col in df.select_dtypes(include=[np.number]).columns:
            mask = df[col].notna() & (df[col] < 0)
            for idx in df.index[mask]:
                invalid_rows.append({
                    "Sheet": sheet, "Excel_Row": int(idx) + 2,
                    "Column": str(col), "Value": df.loc[idx, col],
                    "Issue": "Giá trị số âm; cần kiểm tra quy tắc nghiệp vụ."
                })

        # Cảnh báo ô chỉ chứa khoảng trắng ở cột văn bản
        for col in df.select_dtypes(include=["object", "string"]).columns:
            text_series = df[col].astype("string")
            mask = text_series.notna() & text_series.str.strip().eq("")
            for idx in df.index[mask]:
                invalid_rows.append({
                    "Sheet": sheet, "Excel_Row": int(idx) + 2,
                    "Column": str(col), "Value": repr(df.loc[idx, col]),
                    "Issue": "Chuỗi rỗng hoặc chỉ có khoảng trắng."
                })

    with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
        pd.DataFrame(overview_rows).to_excel(writer, sheet_name="Overview", index=False)
        pd.DataFrame(profile_rows).to_excel(writer, sheet_name="Column_Profile", index=False)
        pd.DataFrame(duplicate_rows).to_excel(writer, sheet_name="Duplicates", index=False)
        pd.DataFrame(outlier_rows, columns=[
            "Sheet", "Excel_Row", "Column", "Value", "IQR_Lower_Bound",
            "IQR_Upper_Bound", "Note"
        ]).to_excel(writer, sheet_name="Outliers_IQR", index=False)
        pd.DataFrame(invalid_rows, columns=[
            "Sheet", "Excel_Row", "Column", "Value", "Issue"
        ]).to_excel(writer, sheet_name="Invalid_Values", index=False)

        # Định dạng báo cáo dễ đọc trong Excel
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        for ws in writer.book.worksheets:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.row_dimensions[1].height = 30
            for col_cells in ws.columns:
                max_len = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells[:200])
                ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(max(max_len + 2, 12), 42)
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)

    print(f"Đã hoàn thành Data Profiling.")
    print(f"File nguồn: {INPUT_FILE}")
    print(f"Báo cáo: {OUTPUT_FILE}")
    for item in overview_rows:
        print(f"- {item['Sheet']}: {item['Rows']} dòng, {item['Columns']} cột, "
              f"{item['Exact_Duplicate_Rows']} dòng trùng hoàn toàn, "
              f"{item['Total_Null_Cells']} ô NULL")

if __name__ == "__main__":
    main()

